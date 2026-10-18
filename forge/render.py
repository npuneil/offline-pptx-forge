"""Spec (JSON) -> native PPTX. Deterministic; the LLM never touches drawing code directly."""
import json
from pathlib import Path

from pptx import Presentation

from . import themes
from .draw import Canvas, E, mix, words
from .infographics import LAYOUTS, MS_COLORS, g, items

W, H = 13.333, 7.5
NON_CANVAS_ROLES = {"ornament"}


class Report:
    def __init__(self):
        self.slides = {}

    def start(self, idx, layout):
        self.slides[idx] = {"layout": layout, "texts": [], "icons": [], "issues": []}

    def text(self, idx, role, s, pt, max_pt, overflow):
        self.slides[idx]["texts"].append({"role": role, "words": words(s), "pt": pt, "max_pt": max_pt,
                                          "overflow": overflow, "text": s[:80]})

    def icon(self, idx, name):
        self.slides[idx]["icons"].append(name)

    def issue(self, idx, kind, detail=""):
        self.slides[idx]["issues"].append({"kind": kind, "detail": str(detail)[:200]})

    def summary(self):
        out = []
        for idx in sorted(self.slides):
            s = self.slides[idx]
            vis = [t for t in s["texts"] if t["role"] not in NON_CANVAS_ROLES]
            out.append({
                "index": idx, "layout": s["layout"],
                "words": sum(t["words"] for t in vis),
                "overflows": sum(1 for t in vis if t["overflow"]),
                "shrunk": sum(1 for t in vis if t["pt"] < t["max_pt"] - 2.5),
                "icons": len(s["icons"]), "issues": s["issues"],
            })
        return out


def _background(slide, t):
    fill = slide.background.fill
    if t.bg2 and t.bg2 != t.bg:
        fill.gradient()
        fill.gradient_angle = 90 if not t.dark else 135
        st = fill.gradient_stops
        st[0].color.rgb = _rgb(t.bg)
        st[0].position = 0.0
        st[1].color.rgb = _rgb(t.bg2)
        st[1].position = 1.0
    else:
        fill.solid()
        fill.fore_color.rgb = _rgb(t.bg)


def _rgb(h):
    from .draw import rgb
    return rgb(h)


def _ms_logo(c, x, y, s, gap):
    for k, col in enumerate(MS_COLORS):
        c.box(x + (k % 2) * (s + gap), y + (k // 2) * (s + gap), s, s, fill=col, shape="rect")


def _ambient(c, title=False):
    t = c.t
    if t.motif == "surface_bloom":
        c.box(9.6 if not title else 7.2, -2.2, 6.5, 6.0, fill=mix(t.primary, t.bg, 0.55), shape="ellipse", alpha=10)
        c.box(11.0 if not title else 9.4, -0.6, 4.5, 4.2, fill=mix(t.accent(2), t.bg, 0.4), shape="ellipse", alpha=9)
    elif t.motif == "ultra_glow":
        c.box(9.8 if not title else 6.8, -3.0, 7.0, 7.0, fill=t.primary, shape="ellipse", alpha=5)
        c.box(11.2 if not title else 9.0, -1.5, 4.5, 4.5, fill=t.accent(1), shape="ellipse", alpha=4)


def chrome(c, s, idx, total, deck):
    t = c.t
    _ambient(c)
    if t.motif == "ultra_glow":
        c.box(0.6, 0.32, 0.75, 0.06, fill=t.primary, shape="rect", glow=(t.primary, 6, 60))
    elif t.motif == "surface_bloom":
        c.box(0.6, 0.32, 0.38, 0.06, fill=t.primary, shape="rect")
    c.text(0.6, 0.45, 12.1, 0.8, g(s, "title"), size=32, font=t.font_head, anchor="m", role="title")
    tk = g(s, "takeaway", "subtitle")
    if tk:
        c.text(0.6, 1.25, 12.1, 0.55, tk, size=17, color=t.subtext, anchor="t", role="takeaway")
    fy = 7.02
    fx = 0.6
    if t.motif == "ms_squares":
        _ms_logo(c, 0.6, fy + 0.07, 0.09, 0.02)
        fx = 0.86
    client = g(deck, "client")
    c.text(fx, fy, 8, 0.3, f"{t.footer}" + (f"  |  {client}" if client else ""), size=10, color=t.muted,
           role="ornament")
    c.text(11.9, fy, 0.83, 0.3, f"{idx + 1}", size=10, color=t.muted, align="r", role="ornament")


def title_slide(c, s, deck):
    t = c.t
    _ambient(c, title=True)
    hero = [h for h in (s.get("hero_icons") or []) if isinstance(h, str)][:4] or ["cloud", "laptop", "shield", "ai"]
    while len(hero) < 4:
        hero.append("idea")
    if t.motif == "ms_squares":
        _ms_logo(c, 0.7, 0.62, 0.17, 0.03)
        c.text(1.15, 0.58, 3, 0.42, "Microsoft", size=19, font=t.font_head, color="737373", anchor="m", role="ornament")
        S, G = 1.75, 0.16
        x0, y0 = 8.75, 1.75
        for k, col in enumerate(MS_COLORS):
            bx, by = x0 + (k % 2) * (S + G), y0 + (k // 2) * (S + G)
            c.box(bx, by, S, S, fill=col, shape="rect")
            c.icon(bx + S * 0.25, by + S * 0.25, S * 0.5, hero[k], "FFFFFF")
    elif t.motif == "surface_bloom":
        c.box(8.6, 1.55, 3.9, 3.9, fill="FFFFFF", shape="ellipse", shadow=True)
        c.icon(9.5, 2.45, 2.1, hero[1], t.primary)
        for k, (bx, by) in enumerate(((8.15, 1.35), (11.75, 2.0), (8.35, 4.7), (11.55, 4.75))):
            col = t.accent(k)
            c.box(bx, by, 0.95, 0.95, fill="FFFFFF", shape="ellipse", shadow=True)
            c.icon(bx + 0.2, by + 0.2, 0.55, hero[(k + 2) % 4] if k else hero[0], col)
    else:
        cx, cy = 10.35, 3.7
        for k, r in enumerate((2.35, 1.85, 1.35)):
            c.box(cx - r, cy - r, 2 * r, 2 * r, line=t.accent(k), line_w=1.25 + k * 0.5, shape="ellipse",
                  glow=(t.accent(k), 8, 30))
        c.box(cx - 0.95, cy - 0.95, 1.9, 1.9, fill=t.card, shape="ellipse", glow=(t.primary, 24, 45))
        c.icon(cx - 0.68, cy - 0.7, 1.36, hero[1], t.primary)
        for k in range(4):
            px, py = [(cx - 2.35, cy), (cx + 2.35, cy), (cx, cy - 2.35), (cx, cy + 2.35)][k]
            c.box(px - 0.38, py - 0.38, 0.76, 0.76, fill=t.card, line=t.accent(k), line_w=1.5, shape="ellipse")
            c.icon(px - 0.23, py - 0.23, 0.46, hero[(k + 2) % 4] if k != 1 else hero[0], t.accent(k))
    kicker = g(s, "kicker", default=f"Prepared for {g(deck, 'client')}" if g(deck, "client") else "")
    if kicker:
        c.text(0.7, 1.85, 7.4, 0.45, kicker, size=15, font=t.font_head, color=t.primary, caps=True, role="takeaway")
    c.text(0.7, 2.35, 7.0, 2.4, g(s, "title", default=g(deck, "title")), size=50, font=t.font_head, anchor="t",
           role="title")
    c.text(0.7, 4.85, 7.2, 1.0, g(s, "subtitle", "takeaway"), size=20, color=t.subtext, role="takeaway")
    foot = g(s, "presenter", default=t.footer)
    c.text(0.7, 6.6, 7.0, 0.35, foot, size=12, color=t.muted, role="ornament")


def section_slide(c, s, deck, idx):
    t = c.t
    _ambient(c, title=True)
    c.text(0.7, 1.6, 6, 2.2, g(s, "number", default=f"{idx:02d}"), size=120, font=t.font_number, color=t.primary,
           role="ornament")
    c.text(0.7, 4.0, 10, 1.2, g(s, "title"), size=44, font=t.font_head, role="title")
    c.text(0.7, 5.2, 10, 0.8, g(s, "takeaway", "subtitle"), size=20, color=t.subtext, role="takeaway")


def render(spec, template, out_path):
    t = themes.get(template)
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(W), E(H)
    blank = prs.slide_layouts[6]
    report = Report()
    slides = spec.get("slides") or []
    for idx, s in enumerate(slides):
        s = s if isinstance(s, dict) else {"layout": "statement", "title": str(s)}
        layout = str(s.get("layout", "statement")).lower().strip()
        slide = prs.slides.add_slide(blank)
        _background(slide, t)
        report.start(idx, layout)
        c = Canvas(slide, t, report, idx)
        try:
            if layout == "title":
                title_slide(c, s, spec)
            elif layout == "section":
                section_slide(c, s, spec, idx)
            elif layout in LAYOUTS:
                chrome(c, s, idx, len(slides), spec)
                LAYOUTS[layout](c, s.get("data") if isinstance(s.get("data"), dict) else s)
            else:
                raise ValueError(f"unknown layout '{layout}'")
        except Exception as exc:  # per-slide fallback keeps the deck whole and records the defect
            for shp in list(slide.shapes):
                shp._element.getparent().remove(shp._element)
            report.start(idx, layout)
            report.issue(idx, "fallback", f"{type(exc).__name__}: {exc}")
            c = Canvas(slide, t, report, idx)
            chrome(c, {"title": g(s, "title", default="(untitled)"), "takeaway": g(s, "takeaway")}, idx, len(slides), spec)
            LAYOUTS["statement"](c, {"icon": "warning", "text": g(s, "takeaway", "title", default="")})
        notes = g(s, "notes", "speaker_notes")
        if notes:
            slide.notes_slide.notes_text_frame.text = str(notes)
    prs.core_properties.title = str(g(spec, "title", default="Offline PPTX Forge deck"))
    prs.core_properties.author = "Offline PPTX Forge"
    prs.core_properties.comments = f"Generated offline; template={template}"
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return {"template": template, "slides": report.summary(), "path": str(out_path)}


def render_file(spec_path, template, out_path):
    spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
    return render(spec, template, out_path)
