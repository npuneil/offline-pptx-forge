"""Native-shape drawing primitives with measured text fitting (PIL metrics of the real Segoe fonts)."""
import math
from functools import lru_cache
from pathlib import Path

from lxml import etree
from PIL import ImageFont
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .icons import ICON_FONT, glyph
from .themes import FONT_FILES

FONTS_DIR = Path(r"C:\Windows\Fonts")
LINE_HEIGHT = 1.30  # Segoe UI ascent+descent ~ 1.33 em
SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE, "rounded": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL,
    "chevron": MSO_SHAPE.CHEVRON, "pentagon": MSO_SHAPE.PENTAGON, "hexagon": MSO_SHAPE.HEXAGON,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE, "arrow": MSO_SHAPE.RIGHT_ARROW, "donut": MSO_SHAPE.DONUT,
    "trapezoid": MSO_SHAPE.TRAPEZOID, "diamond": MSO_SHAPE.DIAMOND, "arc": MSO_SHAPE.BLOCK_ARC,
}
ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def E(inches):
    return Emu(int(round(inches * 914400)))


def rgb(hexstr):
    return RGBColor.from_string(hexstr.strip("#").upper()[:6])


def mix(a, b, t):
    """Blend hex colour a toward b by t (0..1)."""
    a, b = a.strip("#"), b.strip("#")
    ca = [int(a[i:i + 2], 16) for i in (0, 2, 4)]
    cb = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ca, cb))


@lru_cache(maxsize=16)
def _font(name):
    return ImageFont.truetype(str(FONTS_DIR / FONT_FILES.get(name, "segoeui.ttf")), 100)


@lru_cache(maxsize=20000)
def _width100(font_name, s):
    return _font(font_name).getlength(s)


def wrap_lines(s, font_name, pt, width_in):
    """Greedy word wrap using real glyph advances. Returns list of lines."""
    avail = width_in * 72.0
    lines = []
    for para in str(s).split("\n"):
        words = para.split(" ")
        cur = ""
        for w in words:
            cand = (cur + " " + w) if cur else w
            if _width100(font_name, cand) * pt / 100.0 <= avail or not cur:
                cur = cand
            else:
                lines.append(cur)
                cur = w
        lines.append(cur)
    return lines


def fit(s, font_name, w_in, h_in, max_pt, min_pt=9.0):
    """Largest size <= max_pt where wrapped text fits the box. Returns (pt, lines, overflow)."""
    pt = max_pt
    while True:
        lines = wrap_lines(s, font_name, pt, w_in)
        longest = max((_width100(font_name, w) * pt / 100.0 for ln in lines for w in ln.split(" ") if w), default=0)
        fits = len(lines) * pt * LINE_HEIGHT <= h_in * 72.0 + 0.5 and longest <= w_in * 72.0 + 0.5
        if fits:
            return pt, len(lines), False
        if pt <= min_pt:
            return pt, len(lines), True
        pt = max(min_pt, pt - 1.0)


def words(s):
    return len([w for w in str(s or "").replace("\n", " ").split(" ") if w.strip()])


class Canvas:
    """Wraps one slide; records every text element for density and overflow grading."""

    def __init__(self, slide, theme, report, index):
        self.slide = slide
        self.t = theme
        self.report = report
        self.index = index

    # ---- shapes -----------------------------------------------------------------
    def box(self, x, y, w, h, fill=None, line=None, line_w=1.0, shape="rounded", radius=None,
            alpha=None, gradient=None, gradient_angle=90, shadow=False, glow=None, dash=False, rotation=0):
        shp = self.slide.shapes.add_shape(SHAPES.get(shape, MSO_SHAPE.ROUNDED_RECTANGLE), E(x), E(y), E(w), E(h))
        if shape == "rounded":
            r = self.t.radius if radius is None else radius  # absolute corner radius, inches
            shp.adjustments[0] = max(0.0, min(0.5, r / max(0.01, min(w, h))))
        if gradient:
            shp.fill.gradient()
            shp.fill.gradient_angle = gradient_angle
            stops = shp.fill.gradient_stops
            stops[0].color.rgb = rgb(gradient[0])
            stops[0].position = 0.0
            stops[1].color.rgb = rgb(gradient[1])
            stops[1].position = 1.0
            if alpha is not None:
                for gs in shp._element.spPr.iter(qn("a:srgbClr")):
                    _alpha(gs, alpha)
        elif fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
            if alpha is not None:
                _alpha(shp._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr")), alpha)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(line_w)
            if dash:
                shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            shp.line.fill.background()
        if rotation:
            shp.rotation = rotation
        _effects(shp, shadow=shadow and not self.t.dark, glow=glow)
        shp.text_frame.text = ""
        return shp

    def line(self, x1, y1, x2, y2, color=None, width=1.5, arrow=False, dash=False, alpha=None):
        c = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
        c.line.color.rgb = rgb(color or self.t.muted)
        c.line.width = Pt(width)
        if dash:
            c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        ln = c.line._get_or_add_ln()
        if alpha is not None:
            _alpha(ln.find(qn("a:solidFill")).find(qn("a:srgbClr")), alpha)
        if arrow:
            ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
        return c

    # ---- text -------------------------------------------------------------------
    def text(self, x, y, w, h, s, size=18, color=None, font=None, align="l", anchor="t",
             min_size=9, role="body", italic=False, spacing=None, caps=False):
        s = "" if s is None else str(s)
        if caps:
            s = s.upper()
        font = font or self.t.font_body
        pt, nlines, overflow = fit(s, font, w - 0.02, h, size, min_size)
        tb = self.slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = ANCHOR[anchor]
        for i, para in enumerate(s.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN[align]
            if spacing:
                p.line_spacing = spacing
            r = p.add_run()
            r.text = para
            r.font.size = Pt(pt)
            r.font.name = font
            r.font.italic = italic
            r.font.color.rgb = rgb(color or self.t.text)
        self.report.text(self.index, role, s, pt, size, overflow)
        return tb

    def icon(self, x, y, size, name, color=None):
        tb = self.slide.shapes.add_textbox(E(x), E(y), E(size), E(size))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = glyph(name)
        r.font.name = ICON_FONT
        r.font.size = Pt(size * 72 * 0.62)
        r.font.color.rgb = rgb(color or self.t.text)
        self.report.icon(self.index, name)
        return tb

    def badge(self, cx, cy, d, name, fill=None, color=None, ring=None, glow=None):
        fill = fill or self.t.primary
        self.box(cx - d / 2, cy - d / 2, d, d, fill=fill, shape="ellipse", line=ring, line_w=2,
                 glow=glow if self.t.dark else None)
        self.icon(cx - d * 0.36, cy - d * 0.36, d * 0.72, name, color or self.t.on_accent)


def _alpha(clr, pct):
    if clr is None:
        return
    for old in clr.findall(qn("a:alpha")):
        clr.remove(old)
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(max(0, min(100, pct)) * 1000))}))


def _effects(shp, shadow=False, glow=None):
    spPr = shp._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    eff = spPr.makeelement(qn("a:effectLst"), {})
    if glow:
        color, rad_pt, alpha = glow
        g = eff.makeelement(qn("a:glow"), {"rad": str(int(rad_pt * 12700))})
        c = g.makeelement(qn("a:srgbClr"), {"val": color})
        c.append(c.makeelement(qn("a:alpha"), {"val": str(int(alpha * 1000))}))
        g.append(c)
        eff.append(g)
    if shadow:
        s = eff.makeelement(qn("a:outerShdw"), {"blurRad": "228600", "dist": "50800", "dir": "5400000",
                                                "algn": "t", "rotWithShape": "0"})
        c = s.makeelement(qn("a:srgbClr"), {"val": "000000"})
        c.append(c.makeelement(qn("a:alpha"), {"val": "14000"}))
        s.append(c)
        eff.append(s)
    # effectLst must follow ln when present
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        ln.addnext(eff)
    else:
        spPr.append(eff)


def polar(cx, cy, rx, ry, i, n, start=-90):
    a = math.radians(start + 360.0 * i / n)
    return cx + rx * math.cos(a), cy + ry * math.sin(a)
